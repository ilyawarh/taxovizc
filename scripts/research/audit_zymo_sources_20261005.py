"""Extract workbook/reference evidence and verify remote-file checksums, read-only inputs."""
from pathlib import Path
from collections import Counter
import hashlib,json,re,zipfile
import openpyxl,pandas as pd
P=Path(__file__).resolve().parents[2]/'runs/zymo_fecal_20261005'; D=P/'derived'
z=zipfile.ZipFile(P/'references/in2116.16S_210122.zymo.zip')
for suffix in ['Composition_Summary_L7.txt','Heatmap/7.species/taxa_abun_table_L7.txt','Composition_Summary_L6.txt']:
 n=next(n for n in z.namelist() if n.endswith(suffix)); (D/('reference_16S_'+suffix.split('/')[-1])).write_bytes(z.read(n))
z=zipfile.ZipFile(P/'references/in2165.DNA.report.zip')
for rank,num in [('genus','3.Genus'),('species','4.Species')]:
 n=next(n for n in z.namelist() if n.endswith('/All/AbundanceTables/'+num+'/abun_table.tsv')); (D/f'zymo_reference_DNA_{rank}.tsv').write_bytes(z.read(n))
w=openpyxl.load_workbook(P/'raw/Zymo_fecal.xlsx',data_only=True); formulas=openpyxl.load_workbook(P/'raw/Zymo_fecal.xlsx',data_only=False); s=w.worksheets[0]
profile=[{'genus_or_group':s.cell(i,1).value,'species_or_group':s.cell(i,2).value,'percent':s.cell(i,3).value,'cell':f'C{i}'} for i in range(3,357)]
pd.DataFrame(profile).to_csv(D/'workbook_reference_16S.csv',index=False)
ref=pd.read_csv(D/'reference_16S_Composition_Summary_L7.txt',sep='\t',skiprows=1)
a=Counter(round(r['percent'],6) for r in profile); b=Counter(round(x*100,6) for x in ref.iloc[:,1]); assert a==b
(D/'reference_provenance_check.json').write_text(json.dumps({'workbook_rows':len(profile),'manufacturer_rows':len(ref),'abundance_multisets_equal_6_decimal_percent':True,'workbook_percent_sum':sum(r['percent'] for r in profile),'interpretation':'Matching numerical profile plus named leading taxa supports source provenance; does not prove product/lot identity.'},indent=2))
rows=[{'target':w.worksheets[1].cell(i,1).value,'cached_value':w.worksheets[1].cell(i,2).value,'formula':formulas.worksheets[1].cell(i,2).value,'description':w.worksheets[1].cell(i,3).value,'source_cell':f'Лист2!B{i}'} for i in range(2,41)]
pd.DataFrame(rows).to_csv(D/'qpcr_formulas.csv',index=False)
expected={}
for line in (P/'final_remote_manifest.txt').read_text(encoding='utf8').splitlines():
 m=re.match(r'^([0-9a-f]{64})  (.+)$',line)
 if m: expected[m[2]]=m[1]
manifest=[]
for name,sha in expected.items():
 q=P/'raw'/name
 with q.open('rb') as fh: actual=hashlib.file_digest(fh,'sha256').hexdigest()
 assert actual==sha,name
 manifest.append({'path':'raw/'+name,'bytes':q.stat().st_size,'sha256':actual,'remote_match':True})
for q in [P/'raw/Zymo_fecal.xlsx',*sorted((P/'references').glob('*'))]:
 with q.open('rb') as fh: sha=hashlib.file_digest(fh,'sha256').hexdigest()
 manifest.append({'path':str(q.relative_to(P)),'bytes':q.stat().st_size,'sha256':sha})
(D/'artifact_manifest.json').write_text(json.dumps(manifest,indent=2))
print(f'Verified {len(expected)} remote hashes; workbook/manufacturer profile values match ({len(profile)} rows).')